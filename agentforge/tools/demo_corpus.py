"""Build and check the Northwind demo corpus (Release 1 task R1-15).

    uv run python -m tools.demo_corpus build   # demo/build/: corpus files + questionnaires (XLSX)
    uv run python -m tools.demo_corpus check   # consistency + every gold quote is in parsed text

The Markdown policies are the source of truth. The build renders one policy as DOCX and one as
PDF so the demo exercises all three parsers; the check parses the *built* files with the real
sandboxed parser, so gold quotes are verified against exactly what gets indexed.
"""

import json
import shutil
import sys
import textwrap
from pathlib import Path
from typing import Any

from agentforge.knowledge.chunking import normalize
from agentforge.knowledge.parsing import parse_file
from agentforge.knowledge.validation import EXTENSIONS

ROOT = Path(__file__).resolve().parent.parent / "demo"
CORPUS = ROOT / "corpus" / "northwind"
QUESTIONNAIRE = ROOT / "questionnaire" / "northwind-questionnaire.json"
GOLD = ROOT / "eval" / "gold.json"
BUILD = ROOT / "build"
AS_DOCX = {"03-encryption-standard.md"}
AS_PDF = {"06-business-continuity-and-disaster-recovery.md"}

EXPECTED = {"questions": 40, "demo": 12, "unanswerable": 6}


def _load(path: Path) -> Any:
    return json.loads(path.read_text(encoding="utf-8"))


def _markdown_sections(text: str) -> list[tuple[int | None, str]]:
    """(heading level or None, text) for each heading line and paragraph."""
    items: list[tuple[int | None, str]] = []
    for line in text.splitlines():
        stripped = line.strip()
        if not stripped:
            continue
        level = len(stripped) - len(stripped.lstrip("#"))
        if 0 < level <= 6 and stripped[level : level + 1] == " ":
            items.append((level, stripped[level:].strip()))
        else:
            items.append((None, stripped))
    return items


def _to_docx(source: Path, target: Path) -> None:
    import docx

    document = docx.Document()
    for level, text in _markdown_sections(source.read_text(encoding="utf-8")):
        if level is None:
            document.add_paragraph(text)
        else:
            document.add_heading(text, level=min(level, 6))
    document.save(str(target))


def _pdf_escape(text: str) -> str:
    return text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def _to_pdf(source: Path, target: Path) -> None:
    """Plain-text PDF: wrapped lines, a blank line between paragraphs, 46 lines per page."""
    lines: list[str] = []
    for _, text in _markdown_sections(source.read_text(encoding="utf-8")):
        lines += [*textwrap.wrap(text, width=88), ""]
    pages = [lines[i : i + 46] for i in range(0, len(lines), 46)]
    objects: list[bytes] = [b"<< /Type /Catalog /Pages 2 0 R >>"]
    kids = " ".join(f"{3 + 2 * i} 0 R" for i in range(len(pages)))
    objects.append(f"<< /Type /Pages /Kids [{kids}] /Count {len(pages)} >>".encode())
    font = 3 + 2 * len(pages)
    for i, page in enumerate(pages):
        ops = ["BT /F1 10 Tf 14 TL 60 750 Td"]
        ops += [f"({_pdf_escape(line)}) Tj T*" for line in page]
        stream = (" ".join(ops) + " ET").encode("latin-1", "replace")
        objects.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents {4 + 2 * i} 0 R "
            f"/Resources << /Font << /F1 {font} 0 R >> >> >>".encode()
        )
        objects.append(b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream")
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    out = bytearray(b"%PDF-1.4\n")
    offsets: list[int] = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n".encode() + body + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    for offset in offsets:
        out += f"{offset:010d} 00000 n \n".encode()
    out += (
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    )
    target.write_bytes(bytes(out))


def _write_xlsx(questions: list[dict[str, Any]], target: Path, title: str) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font

    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None  # noqa: S101 - a new workbook always has one sheet
    sheet.title = "Questionnaire"
    sheet.append(["ID", "Section", "Question", "Answer", "Notes"])
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for q in questions:
        sheet.append([q["id"], q["section"], q["question"], "", ""])
    for column, width in zip("ABCDE", (8, 22, 70, 60, 30), strict=True):
        sheet.column_dimensions[column].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    workbook.properties.title = title
    workbook.save(str(target))


def built_document_name(markdown_name: str) -> str:
    if markdown_name in AS_DOCX:
        return markdown_name.removesuffix(".md") + ".docx"
    if markdown_name in AS_PDF:
        return markdown_name.removesuffix(".md") + ".pdf"
    return markdown_name


def build() -> Path:
    if BUILD.exists():
        shutil.rmtree(BUILD)
    corpus = BUILD / "corpus"
    corpus.mkdir(parents=True)
    for source in sorted((CORPUS / "policies").glob("*.md")) + sorted(
        (CORPUS / "adversarial").glob("*.md")
    ):
        target = corpus / built_document_name(source.name)
        if source.name in AS_DOCX:
            _to_docx(source, target)
        elif source.name in AS_PDF:
            _to_pdf(source, target)
        else:
            shutil.copyfile(source, target)
    questionnaire = _load(QUESTIONNAIRE)
    questions = questionnaire["questions"]
    _write_xlsx(questions, BUILD / "northwind-questionnaire.xlsx", questionnaire["title"])
    demo = [q for q in questions if q["demo"]]
    _write_xlsx(
        demo, BUILD / "northwind-questionnaire-demo12.xlsx", questionnaire["title"] + " (demo)"
    )
    return BUILD


def check() -> list[str]:
    problems: list[str] = []
    questions = _load(QUESTIONNAIRE)["questions"]
    answers = {a["id"]: a for a in _load(GOLD)["answers"]}
    ids = [q["id"] for q in questions]
    if len(ids) != EXPECTED["questions"] or len(set(ids)) != len(ids):
        problems.append(f"expected {EXPECTED['questions']} unique questions, found {len(set(ids))}")
    if set(ids) != set(answers):
        problems.append(f"questionnaire and gold IDs differ: {sorted(set(ids) ^ set(answers))}")
    demo = [q["id"] for q in questions if q["demo"]]
    unanswerable = [i for i, a in answers.items() if not a["answerable"]]
    if len(demo) != EXPECTED["demo"]:
        problems.append(f"expected {EXPECTED['demo']} demo questions, found {len(demo)}")
    if len(unanswerable) != EXPECTED["unanswerable"]:
        problems.append(
            f"expected {EXPECTED['unanswerable']} unanswerable, found {len(unanswerable)}"
        )
    if not set(demo) & set(unanswerable) or not any(
        answers[i].get("injection_target") for i in demo
    ):
        problems.append("demo subset must include unanswerable and injection-target questions")

    texts: dict[str, str] = {}
    corpus = build() / "corpus"
    for path in sorted(corpus.iterdir()):
        media_type = EXTENSIONS[path.suffix.lower()]
        texts[path.name] = normalize(parse_file(path, media_type)).text
    for answer in answers.values():
        if answer["answerable"] and not answer["evidence"]:
            problems.append(f"{answer['id']}: answerable but has no evidence")
        for evidence in answer["evidence"]:
            name = built_document_name(evidence["document"])
            text = texts.get(name)
            if text is None:
                problems.append(f"{answer['id']}: unknown document {evidence['document']}")
            elif evidence["quote"] not in text:
                problems.append(
                    f"{answer['id']}: quote not found in parsed {name}: {evidence['quote'][:60]}"
                )
    return problems


def main() -> int:
    command = sys.argv[1] if len(sys.argv) > 1 else "check"
    if command == "build":
        print(f"built {build()}")
        return 0
    problems = check()
    for problem in problems:
        print(f"FAIL {problem}")
    files = len(list((BUILD / "corpus").iterdir()))
    print(f"demo-check: {files} corpus files, 40 questions, {len(problems)} problem(s)")
    return 1 if problems else 0


if __name__ == "__main__":
    sys.exit(main())
